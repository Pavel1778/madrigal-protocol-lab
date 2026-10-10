"""Build the defence deck as a real, editable PPTX from our own code.

The Slidev export shipped a PPTX that is only pictures: no text runs, no
transitions, no embedded fonts. This script rebuilds the same deck with
python-pptx as native editable text, the task palette, a 16:9 grid, entrance
animations and slide transitions.

    python -m presentation.scripts.build_pptx

Layout is fixed in inches; 60 px at 96 dpi is 0.625 in.
"""

from __future__ import annotations

import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
PRES = ROOT / "presentation"
OUT = PRES / "slides.pptx"

# --- palette (task block 2) -------------------------------------------------
BG = RGBColor(0x0E, 0x10, 0x11)
PANEL = RGBColor(0x1A, 0x1C, 0x1E)
TEXT = RGBColor(0xF0, 0xF0, 0xF0)
SECONDARY = RGBColor(0xB0, 0xB0, 0xB0)
CAPTION = RGBColor(0x9A, 0x9A, 0x9A)
HIGHLIGHT = RGBColor(0xA2, 0x39, 0x1D)
ACCENT_TEXT = RGBColor(0xE0, 0x60, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

F_HEAD = "Manrope"       # headings: Inter/Manrope Bold per task
F_BODY = "Inter"         # body: Inter Regular
F_MONO = "JetBrains Mono"  # code

MARGIN = Inches(0.625)
FOOTER_H = Inches(0.5)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
BODY_TOP = Inches(1.7)
GITHUB = "https://github.com/Pavel1778/madrigal-protocol-lab"

_TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_runs(paragraph, text, size, color, bold=False, font=F_BODY):
    """Write text into a paragraph, honouring **bold** and `code` spans."""
    for part in _TOKEN.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        run.font.size = Pt(size)
        run.font.name = font
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.font.bold = True
            run.font.color.rgb = WHITE
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.font.name = F_MONO
            run.font.size = Pt(size - 2)
            run.font.color.rgb = ACCENT_TEXT
        else:
            run.text = part
            run.font.bold = bold
            run.font.color.rgb = color


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tb, tf


def solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    solid(s, color)
    return s


def bg(slide):
    rect(slide, 0, 0, SLIDE_W, SLIDE_H, BG)
    rect(slide, 0, 0, Inches(0.06), SLIDE_H, HIGHLIGHT)


def footer(slide, page, total=22):
    tb, tf = textbox(slide, MARGIN, SLIDE_H - FOOTER_H,
                     SLIDE_W - 2 * MARGIN, Inches(0.3))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = f"Команда HTTP 418 · UMIRHack 2026 · {page} / {total}"
    run.font.size = Pt(10)
    run.font.name = F_BODY
    run.font.color.rgb = CAPTION


def heading(slide, text, size=32):
    tb, tf = textbox(slide, MARGIN, Inches(0.55),
                     SLIDE_W - 2 * MARGIN, Inches(1.0))
    add_runs(tf.paragraphs[0], text, size, WHITE, bold=True, font=F_HEAD)
    return tb


def caption(slide, text, y=Inches(6.55)):
    tb, tf = textbox(slide, MARGIN, y, SLIDE_W - 2 * MARGIN, Inches(0.5))
    add_runs(tf.paragraphs[0], text, 11, CAPTION)
    return tb


def bullets(slide, items, x, y, w, h, size=17):
    tb, tf = textbox(slide, x, y, w, h)
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(10)
        add_runs(p, "•  " + it, size, TEXT)
    return tb


def picture(slide, rel, x, y, w, h=None):
    path = PRES / rel.lstrip("/")
    pic = slide.shapes.add_picture(str(path), x, y, width=w)
    if h is not None:
        pic.height = h
    return pic


def fit_picture(slide, rel, x, y, max_w, max_h):
    """Place a picture scaled to fit inside max_w x max_h, keeping aspect."""
    from PIL import Image

    path = PRES / rel.lstrip("/")
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w = int(iw * scale)
    h = int(ih * scale)
    return slide.shapes.add_picture(str(path), x, y, width=w, height=h)


# --- animations / transitions ----------------------------------------------
NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def _fade(spid: int, cid: int) -> str:
    return f"""
<p:par><p:cTn id="{cid}" presetID="10" presetClass="entr" presetSubtype="0"
 fill="hold" grpId="0" nodeType="clickEffect">
 <p:stCondLst><p:cond delay="0"/></p:stCondLst>
 <p:childTnLst>
  <p:set>
   <p:cBhvr><p:cTn id="{cid+1}" dur="1" fill="hold">
    <p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>
    <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
    <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
   </p:cBhvr>
   <p:to><p:strVal val="visible"/></p:to>
  </p:set>
  <p:anim calcmode="lin" valueType="num">
   <p:cBhvr additive="base">
    <p:cTn id="{cid+2}" dur="500" fill="hold"/>
    <p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>
    <p:attrNameLst><p:attrName>style.opacity</p:attrName></p:attrNameLst>
   </p:cBhvr>
   <p:tavLst>
    <p:tav tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav>
    <p:tav tm="100000"><p:val><p:fltVal val="1"/></p:val></p:tav>
   </p:tavLst>
  </p:anim>
 </p:childTnLst>
</p:cTn></p:par>"""


def add_entrance(slide, shape_ids: list[int]) -> None:
    """Append a click-triggered fade entrance for each shape id."""
    if not shape_ids:
        return
    cid = 3
    pars = []
    for spid in shape_ids:
        pars.append(_fade(spid, cid))
        cid += 3
    xml = f"""<p:timing xmlns:p="{NS['p']}" xmlns:a="{NS['a']}">
 <p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
  <p:childTnLst><p:seq concurrent="1" nextAc="seek">
   <p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>
    {''.join(pars)}
   </p:childTnLst></p:cTn>
  </p:seq></p:childTnLst>
 </p:cTn></p:par></p:tnLst>
</p:timing>"""
    slide._element.append(etree.fromstring(xml))


def add_transition(slide, kind="fade") -> None:
    xml = f"""<p:transition xmlns:p="{NS['p']}" spd="med" {kind}="1"/>"""
    slide._element.append(etree.fromstring(xml))


def hyperlink(run, url):
    run.hyperlink.address = url


# --- slide builders ---------------------------------------------------------
def build_slide(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    kind = spec["kind"]
    entrance: list[int] = []

    if kind == "cover":
        tb, tf = textbox(slide, MARGIN, Inches(2.2),
                         SLIDE_W - 2 * MARGIN, Inches(1.3))
        add_runs(tf.paragraphs[0], spec["title"], 40, WHITE, font=F_HEAD)
        entrance.append(tb.shape_id)
        tb2, tf2 = textbox(slide, MARGIN, Inches(3.5),
                           SLIDE_W - 2 * MARGIN, Inches(1.0))
        add_runs(tf2.paragraphs[0], spec["subtitle"], 20, SECONDARY)
        tb3, tf3 = textbox(slide, MARGIN, Inches(5.0),
                           SLIDE_W - 2 * MARGIN, Inches(1.2))
        add_runs(tf3.paragraphs[0], spec["meta"], 15, SECONDARY)
        picture(slide, "public/logo.png", SLIDE_W - Inches(1.5), Inches(0.6),
                Inches(0.85), Inches(0.85))

    elif kind == "team":
        picture(slide, "public/team-avatar.png", MARGIN, Inches(2.0),
                Inches(2.6), Inches(2.6))
        heading(slide, spec["title"])
        bullets(slide, spec["bullets"], Inches(3.7), Inches(1.7),
                SLIDE_W - Inches(4.4), Inches(3.5), size=18)
        caption(slide, GITHUB, y=Inches(5.7))

    elif kind == "two_col":
        heading(slide, spec["title"])
        bullets(slide, spec["left"], MARGIN, BODY_TOP,
                Inches(5.9), Inches(4.3), size=17)
        x2 = MARGIN + Inches(6.2)
        w2 = SLIDE_W - x2 - MARGIN
        if spec.get("right_image"):
            fit_picture(slide, spec["right_image"], x2, BODY_TOP, w2,
                        Inches(4.5))
        else:
            bullets(slide, spec["right"], x2, BODY_TOP, w2, Inches(4.3), size=17)
        if spec.get("caption"):
            caption(slide, spec["caption"])

    elif kind == "killer":
        heading(slide, spec["title"])
        gap = Inches(0.3)
        cw = (SLIDE_W - 2 * MARGIN - 2 * gap) / 3
        for i, (tag, body) in enumerate(spec["cards"]):
            x = MARGIN + i * (cw + gap)
            y = Inches(2.1)
            card = rect(slide, x, y, cw, Inches(2.6), PANEL)
            rect(slide, x, y, Inches(0.05), Inches(2.6), HIGHLIGHT)
            tb, tf = textbox(slide, x + Inches(0.2), y + Inches(0.2),
                             cw - Inches(0.4), Inches(2.2))
            add_runs(tf.paragraphs[0], tag, 13, ACCENT_TEXT, font=F_MONO)
            add_runs(tf.add_paragraph(), body, 15, TEXT)
            entrance.append(card.shape_id)
        if spec.get("caption"):
            caption(slide, spec["caption"])

    elif kind == "metrics":
        heading(slide, spec["title"])
        gap = Inches(0.25)
        kw = (Inches(7.2) - gap) / 2
        for i, (num, lab) in enumerate(spec["kpis"]):
            x = MARGIN + (i % 2) * (kw + gap)
            y = BODY_TOP + (i // 2) * Inches(1.15)
            tb, tf = textbox(slide, x, y, kw, Inches(1.0))
            add_runs(tf.paragraphs[0], num, 34, WHITE, font=F_HEAD)
            add_runs(tf.add_paragraph(), lab, 13, SECONDARY)
        x2 = MARGIN + Inches(7.5)
        w2 = SLIDE_W - x2 - MARGIN
        if spec.get("right_image"):
            fit_picture(slide, spec["right_image"], x2, BODY_TOP, w2,
                        Inches(4.4))
        if spec.get("caption"):
            caption(slide, spec["caption"])

    elif kind == "image_full":
        heading(slide, spec["title"])
        fit_picture(slide, spec["image"], MARGIN, BODY_TOP,
                    SLIDE_W - 2 * MARGIN, Inches(4.4))
        if spec.get("caption"):
            caption(slide, spec["caption"])

    elif kind == "divider":
        tb, tf = textbox(slide, MARGIN, Inches(2.4),
                         SLIDE_W - 2 * MARGIN, Inches(1.4))
        add_runs(tf.paragraphs[0], spec["index"], 60, ACCENT_TEXT, font=F_HEAD)
        entrance.append(tb.shape_id)
        tb2, tf2 = textbox(slide, MARGIN, Inches(3.5),
                           SLIDE_W - 2 * MARGIN, Inches(1.0))
        add_runs(tf2.paragraphs[0], spec["title"], 36, WHITE, font=F_HEAD)
        rect(slide, MARGIN, Inches(4.7), Inches(4.0), Inches(0.08), HIGHLIGHT)

    elif kind == "ending":
        tb, tf = textbox(slide, MARGIN, Inches(3.0),
                         SLIDE_W - 2 * MARGIN, Inches(1.2))
        add_runs(tf.paragraphs[0], spec["title"], 40, WHITE, font=F_HEAD)
        tb2, tf2 = textbox(slide, MARGIN, Inches(4.2),
                           SLIDE_W - 2 * MARGIN, Inches(0.8))
        add_runs(tf2.paragraphs[0], spec["subtitle"], 18, SECONDARY)
        tb3, tf3 = textbox(slide, MARGIN, Inches(5.0),
                           SLIDE_W - 2 * MARGIN, Inches(0.6))
        run = tf3.paragraphs[0].add_run()
        run.text = GITHUB
        run.font.size = Pt(14)
        run.font.color.rgb = ACCENT_TEXT
        hyperlink(run, GITHUB)

    elif kind == "links":
        heading(slide, spec["title"])
        bullets(slide, spec["left"], MARGIN, BODY_TOP, Inches(6.0),
                Inches(4.3), size=17)
        x2 = MARGIN + Inches(7.0)
        w2 = SLIDE_W - x2 - MARGIN
        picture(slide, "public/qr_repo.png", x2 + (w2 - Inches(2.4)) / 2,
                BODY_TOP, Inches(2.4), Inches(2.4))
        tb, tf = textbox(slide, x2, BODY_TOP + Inches(2.6), w2, Inches(0.6))
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = GITHUB
        run.font.size = Pt(11)
        run.font.color.rgb = CAPTION
        hyperlink(run, GITHUB)

    footer(slide, spec["page"])
    if spec.get("notes"):
        slide.notes_slide.notes_text_frame.text = spec["notes"]
    return slide, entrance


# --- content (mirrors slides.md, 22 slides) --------------------------------
def deck() -> list[dict]:
    return [
        dict(kind="cover", page=1, title="Лаборатория протоколов",
             subtitle="Исследование и описание неизвестного сетевого протокола поверх TCP",
             meta="Команда HTTP 418 · UMIRHack 2026 · кейс #02 · Мадригал\n"
                  "Python 3.12 · PySide6 · dpkt · pytest · Docker",
             notes="Восстановление недокументированного бинарного протокола поверх TCP."),
        dict(kind="team", page=2, title="Состав команды",
             bullets=["**Сабадаш Павел** — архитектура, capture, project, report, координация, backend, GUI · @Pasha1778",
                      "**Предков Никита** — backend, правила, тесты · @petruan"],
             notes="Команда HTTP 418: Сабадаш Павел, Предков Никита."),
        dict(kind="two_col", page=3,
             title="Задача: восстановить протокол, а не угадать его",
             left=["**Дано:** захваты (pcap/pcapng) и журнал действий устройства. Спецификации нет.",
                   "**Нужно:** проверяемая интерпретация структуры сообщений.",
                   "**Отличие:** парсер даёт один ответ; здесь видно, где он верен, а где перестаёт работать."],
             right=["Три вопроса держат доклад: что наблюдалось · где правило работает · почему это чтение.",
                    "Выход — переносимая интерпретация с контрпримерами, привязанными к байтам."],
             caption="Вход: байты + журнал. Выход: интерпретация с контрпримерами."),
        dict(kind="divider", page=4, index="01", title="Подход", notes="Раздел 01."),
        dict(kind="image_full", page=5,
             title="Три модуля, связанные JSON-контрактами",
             image="assets/architecture.png",
             caption="capture · protocol + hypothesis · ui. Контракты зафиксированы JSON Schema draft 2020-12."),
        dict(kind="divider", page=6, index="02", title="Реализация", notes="Раздел 02."),
        dict(kind="two_col", page=7,
             title="Захват потока: границы TCP ≠ границы сообщений",
             left=["Сессия — экземпляр соединения, а не пара адресов.",
                   "Байты восстанавливаются по номерам последовательности; повторная передача даёт происхождение, а не лишние байты.",
                   "Пропуск не заполняется нулями — становится диагностикой.",
                   "Клик по байту показывает пакет, номер последовательности и время."],
             right_image="screenshots/01_main.png",
             caption="Реальное окно: сессии, байты, интерпретация, происхождение, диагностика."),
        dict(kind="two_col", page=8,
             title="Правила и фрейминг: одно описание на весь поток",
             left=["Фрейминг: по длине, фиксированный размер, по маркерам, вручную.",
                   "Типы полей: uint · bytes · enum · string · checksum · computed.",
                   "Поле с пометкой `hypothesis` — предположение о смысле, а не факт.",
                   "Сообщения не перечисляются вручную: правило описывает весь поток."],
             right=["Пример правила:",
                    "framing length_prefixed, length_offset 2, length_size 2, byte_order big",
                    "fields: command offset 0, uint8, hypothesis, expected [1]"],
             caption="Фрагмент examples/corpus_rule_v1.json."),
        dict(kind="two_col", page=9,
             title="Гипотезы и контрпримеры: совпадение — не доказательство",
             left=["Шесть статусов: matched · mismatched · ambiguous · incomplete · uncovered · not_applicable.",
                   "`mismatched` — контрпример, привязанный к байтам и пакету-источнику.",
                   "Правило проверяется на всём корпусе.",
                   "Контрпримеры сохраняются, а не отбрасываются."],
             right_image="screenshots/05_validation.png"),
        dict(kind="killer", page=10, title="Функционал: три киллер-фичи",
             cards=[("R1 · источник", "Каждое утверждение о протоколе привязано к конкретному байту, пакету и времени."),
                    ("R4 · проверка", "Контрпримеры сохраняются, а не отбрасываются. Совпадение ≠ доказательство."),
                    ("R6 · перенос", "Правило для одного захвата применяется к новым данным без ручного описания.")],
             caption="Происхождение, опровержимость, переносимость."),
        dict(kind="divider", page=11, index="03", title="Демонстрация", notes="Раздел 03."),
        dict(kind="image_full", page=12, title="Цикл на корпусе: шесть шагов",
             image="assets/research_timeline.png",
             caption="Наблюдение → гипотеза → проверка → контрпример → уточнение → повторная проверка."),
        dict(kind="two_col", page=13,
             title="Проверка и контрпримеры: точность 0.43 → 1.00",
             left=["Чтение, объявленное правилом v1, даёт 0.43 (60 support, 80 contradict).",
                   "После уточнения v2: 140 matched, 0 контрпримеров.",
                   "matched 60 → 140, mismatched 80 → 0."],
             right_image="assets/precision_v1_v2.png",
             caption="Источник: docs/METRICS.md."),
        dict(kind="two_col", page=14,
             title="Сравнение версий: изменён один список, структура та же",
             left=["Изменён только список допустимых команд: [1] → [1, 2, 3].",
                   "Фрейминг, число полей, смещения и типы идентичны.",
                   "Старый результат хранится со своей версией и помечается outdated."],
             right_image="screenshots/08_diff.png"),
        dict(kind="two_col", page=15,
             title="Контрпример как ценность: байт, который уточнил правило",
             left=["session / direction: s2 / A_to_B",
                   "message offset 0, length 7",
                   "bytes 02 00 00 03 13 00 55",
                   "reason value 2 not in expected [1]",
                   "packet 128 / seq 20001"],
             right=["Уточнение правила:",
                    "`expected`: [1] → [1, 2, 3]",
                    "Один байт, привязанный к пакету-источнику, превратил 80 контрпримеров в 0 (v1 → v2)."],
             caption="Конкретный mismatched-байт, привязанный к пакету."),
        dict(kind="metrics", page=16, title="Метрики: сводка чисел",
             kpis=[("140", "matched, rule v2"), ("0", "counterexamples, v2"),
                   ("10/10", "потоков разобрано"), ("280", "сообщений в захвате 01"),
                   ("40/40", "перенос на захват 02"), ("5/5", "CI зелёный")],
             right_image="assets/coverage_donut.png",
             caption="Источник: docs/METRICS.md, docs/AUDIT.md."),
        dict(kind="two_col", page=17,
             title="Расширения: веб-интерфейс на localhost поверх того же ядра",
             left=["Ядро отделено от GUI контрактами JSON — веб тонкая обёртка, не дубль.",
                   "Те же src.capture и src.protocol, своей логики нет.",
                   "Ставится отдельно: pip install -e \".[web]\"."],
             right_image="public/web.png",
             caption="FastAPI + HTMX, 127.0.0.1, extra .[web]."),
        dict(kind="divider", page=18, index="04", title="Итог", notes="Раздел 04."),
        dict(kind="two_col", page=19, title="Ограничения: честный список",
             left=["Смысл байта флагов не установлен: в корпусе он постоянен.",
                   "Поле значения остаётся гипотезой.",
                   "Ответы разбираются по границам, но не по полям."],
             right=["Альтернативные чтения берутся из заданного набора.",
                    "Корпус синтетический — протокол может не совпадать с устройством.",
                    "Интерпретация верна для проверенных захватов и молчит вне них."]),
        dict(kind="two_col", page=20, title="Направления развития",
             left=["IPv6 и IP-фрагментация.",
                   "Live capture — разбор по мере поступления."],
             right=["Полный автомат состояний — связывать запрос и ответ по транзакции.",
                    "Реальные данные от производителя."],
             caption="Ни один пункт не входит в проверенную область: это расширения, не обещания."),
        dict(kind="links", page=21, title="Итог и ссылки",
             left=["Восстановлена раскладка: код команды, флаги, длина (uint16, BE), полезная нагрузка.",
                   "Команда принимает 1, 2, 3; смещение 4 — идентификатор параметра.",
                   "Правило переносится на второй захват и не переносится на посторонний.",
                   "Каждое утверждение подкреплено байтами, контрпримером или журналом."]),
        dict(kind="ending", page=22, title="Вопросы",
             subtitle="Команда HTTP 418 · UMIRHack 2026",
             notes="QR ведёт на репозиторий; отчёт REPORT.md."),
    ]


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    for spec in deck():
        slide, entrance = build_slide(prs, spec)
        if entrance:
            add_entrance(slide, entrance)
        if spec["kind"] == "divider":
            add_transition(slide, "push")
        elif spec["page"] in {4, 6, 11, 18}:
            add_transition(slide, "fade")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
